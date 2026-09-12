import json
from openpyxl import Workbook
from openpyxl.styles import PatternFill

json_path = r'C:\Users\UNIONIT\Desktop\ลองทำ\checklist_dump.json'
out_path = r'C:\Users\UNIONIT\Downloads\FTSE_Update_Highlighted_v2.xlsx'

with open(json_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

# Strict list of exactly what we found on the website.
# If an indicator matches any of these exactly (or contains these keywords definitively), it is GREEN.
# Otherwise, it defaults to RED (Gap).
strict_met_indicators = [
    # E - Environment
    "การเปิดเผย GHG Scope 1",
    "การเปิดเผย GHG Scope 2",
    "การบริหารพลังงานและแผนลดการใช้พลังงาน",
    "การลงทุนในเทคโนโลยีประหยัดพลังงาน",
    "นโยบายป้องกันมลพิษ",
    "การจัดทำระบบ ISO 14001",
    "การเปิดเผยปริมาณของเสีย",
    "การรีไซเคิล",
    "การจัดการน้ำเสีย",
    "การจัดการกากอุตสาหกรรม",
    "ลงทุนระบบบำบัดน้ำเสีย",
    "Circular Economy",
    "ISO 14001",
    
    # S - Social
    "นโยบายแรงงาน",
    "ไม่ใช้แรงงานเด็ก",
    "ไม่ใช้แรงงานบังคับ",
    "ไม่เลือกปฏิบัติ",
    "ชั่วโมงฝึกอบรม",
    "อาชีวอนามัย",
    "ISO 45001",
    "คณะกรรมการความปลอดภัย",
    "ฝึกอบรมความปลอดภัย",
    "LTIFR",
    "TRIR",
    "อุบัติเหตุ",
    "ผู้เสียชีวิต",
    "ISO 9001",
    "ข้อมูลส่วนบุคคล",
    "ความหลากหลาย",
    "สัดส่วนผู้หญิง",
    "สวัสดิการ",
    "พัฒนาชุมชน",
    "PDPA",
    
    # G - Governance
    "โครงสร้างคณะกรรมการชัดเจน",
    "สัดส่วนกรรมการอิสระ",
    "สัดส่วนกรรมการหญิง",
    "ความหลากหลายของคณะกรรมการ",
    "ความขัดแย้งทางผลประโยชน์",
    "จรรยาบรรณธุรกิจ",
    "ERM",
    "คณะกรรมการบริหารความเสี่ยง",
    "บริหารความเสี่ยงอัตราแลกเปลี่ยน",
    "บริหารความเสี่ยงซัพพลายเชน",
    "COSO",
    "ตรวจสอบภายใน",
    "ต่อต้านคอร์รัปชัน",
    "Whistleblower",
    "Data Privacy"
]

# Specifically flag some things as RED to override any accidental keyword overlap
strict_red_indicators = [
    "Scope 3", "TCFD", "ISSB", "Net Zero", "SBTi", "RE100", 
    "ภาษี", "Tax", "Due Diligence", "NPS", "ความหลากหลายทางชีวภาพ", "Biodiversity",
    "ค่าปรับ", "คดี"
]

def evaluate(indicator):
    ind = str(indicator).lower()
    
    # Check red overrides first
    for red in strict_red_indicators:
        if red.lower() in ind:
            return False
            
    # Check if it's in the met list
    for green in strict_met_indicators:
        if green.lower() in ind:
            return True
            
    return False

wb = Workbook()
wb.remove(wb.active) # Remove default sheet

green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

for sheet_name, categories in data.items():
    ws = wb.create_sheet(title=sheet_name)
    ws.append(["Category", "Indicator", "Status"])
    
    for category, indicators in categories.items():
        for indicator in indicators:
            met = evaluate(indicator)
            row = [category, indicator, "มี (Met)" if met else "ไม่มี (Gap)"]
            ws.append(row)
            
            row_idx = ws.max_row
            fill_color = green_fill if met else red_fill
            for col_idx in range(1, 4):
                ws.cell(row=row_idx, column=col_idx).fill = fill_color
                
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 80
    ws.column_dimensions['C'].width = 15

wb.save(out_path)
print(f"Saved strict highlighted Excel to {out_path}")
