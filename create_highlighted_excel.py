import json
from openpyxl import Workbook
from openpyxl.styles import PatternFill

json_path = r'C:\Users\UNIONIT\Desktop\ลองทำ\checklist_dump.json'
out_path = r'C:\Users\UNIONIT\Downloads\FTSE_Update_Highlighted.xlsx'

with open(json_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

manual_reds = [
    'scope 3', 'tcfd', 'issb', 'net zero', 'sbti', 're100', 'cdp', 
    'tax', 'ภาษี', 'whistleblow', 'แจ้งเบาะแส', 'supplier code', 
    'due diligence', 'nps', 'ความผูกพันพนักงาน', 'water stress',
    'biodiversity', 'ความหลากหลายทางชีวภาพ'
]

manual_greens = [
    'การลดก๊าซเรือนกระจก', 'การประหยัดพลังงาน', 'การจัดการของเสีย',
    'การใช้น้ำ', 'ความปลอดภัย', 'ชุมชน', 'โครงสร้างคณะกรรมการ', 
    'การประเมินความเสี่ยง', 'ต่อต้านคอร์รัปชัน', 'ความขัดแย้งทางผลประโยชน์',
    'scope 1', 'scope 2', 'iso 14001', 'coso', 'กรรมการอิสระ', 
    'ความหลากหลาย', 'iso 9001', 'แรงงานเด็ก', 'บังคับ', 'pdpa'
]

def evaluate(indicator):
    ind = str(indicator).lower()
    for red in manual_reds:
        if red in ind:
            return False
    for green in manual_greens:
        if green in ind:
            return True
    return False

wb = Workbook()
wb.remove(wb.active) # Remove default sheet

green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

for sheet_name, categories in data.items():
    ws = wb.create_sheet(title=sheet_name)
    ws.append(["Category", "Indicator", "Status"])
    
    for category, indicators in categories.items():
        for indicator in indicators:
            met = evaluate(indicator)
            row = [category, indicator, "มี (Met)" if met else "ไม่มี (Gap)"]
            ws.append(row)
            
            # Apply fill
            row_idx = ws.max_row
            fill_color = green_fill if met else red_fill
            for col_idx in range(1, 4):
                ws.cell(row=row_idx, column=col_idx).fill = fill_color
                
    # Adjust column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 80
    ws.column_dimensions['C'].width = 15

wb.save(out_path)
print(f"Saved generated highlighted Excel to {out_path}")
